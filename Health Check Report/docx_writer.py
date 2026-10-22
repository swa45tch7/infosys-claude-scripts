#!/usr/bin/env python3
"""Word (.docx) rendering of the health check assessment.

Layout, colours and sizes follow the Infosys LM Envision health check
deliverable: A4, Arial, navy table headers, RAG-shaded status cells, callout
boxes and KPI tiles. With a template, the template's styles, page setup and
header/footer are kept and only the body is replaced.
"""

from __future__ import annotations

import re
from datetime import timezone
from pathlib import Path
from typing import Any, Callable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor, Twips

from assessment import Assessment

NAVY = "071C4F"
BLUE = "006FF2"
INK = "1E2430"
MUTED = "586273"
HEADING2 = "0D2A6B"
BORDER = "D7DEE8"
PALE = "F5F9FF"
PALE_BLUE = "EAF3FF"
PALE_GREY = "F3F5F7"
WHITE = "FFFFFF"
GREEN_TEXT = "2E7D32"
AMBER_TEXT = "A15C00"
AMBER_FILL = "FFF3D6"
RED_TEXT = "B3261E"

STATUS_COLOURS = {
    "GREEN": ("E7F4E8", GREEN_TEXT),
    "AMBER": (AMBER_FILL, AMBER_TEXT),
    "RED": ("FCE8E6", RED_TEXT),
    "UNKNOWN": ("ECEFF1", "5F6368"),
    "STRONG": ("E7F4E8", GREEN_TEXT),
    "DEVELOPING": (AMBER_FILL, AMBER_TEXT),
    "AT RISK": ("FCE8E6", RED_TEXT),
}
TABLE_WIDTH = 9921
TILE_WIDTH = 3305
FIGURE_WIDTH = Emu(6263999)
MARKUP = re.compile(r"(\[Consultant to [^\]]*\]|Data unavailable: [^.\[]*)")
TCPR_ORDER = (
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
)

SectionProgress = Callable[[str, bool, str], None]


def status_colours(label: str) -> tuple[str, str]:
    """Fill from the first part, text colour from the last ("GREEN / AMBER")."""
    parts = [p.strip().upper() for p in label.split("/")]
    fill = STATUS_COLOURS.get(parts[0], STATUS_COLOURS["UNKNOWN"])[0]
    colour = STATUS_COLOURS.get(parts[-1], STATUS_COLOURS["UNKNOWN"])[1]
    return fill, colour


# --------------------------------------------------------------------------- #
# Low-level helpers
# --------------------------------------------------------------------------- #
def _run(paragraph, text: str, size: float | None = None, colour: str | None = None,
         bold: bool | None = None, italic: bool | None = None, superscript: bool = False):
    run = paragraph.add_run(text)
    font = run.font
    if size:
        font.size = Pt(size)
    if colour:
        font.color.rgb = RGBColor.from_string(colour)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if superscript:
        font.superscript = True
    return run


def _text(paragraph, text: str, size: float | None = None, colour: str = INK, bold: bool = False) -> None:
    """Write text, highlighting consultant placeholders and unavailable-data notes."""
    for piece in MARKUP.split(text):
        if not piece:
            continue
        if piece.startswith("[Consultant to"):
            _run(paragraph, piece, size, AMBER_TEXT, bold, italic=True)
        elif piece.startswith("Data unavailable:"):
            _run(paragraph, piece, size, RED_TEXT, bold, italic=True)
        else:
            _run(paragraph, piece, size, colour, bold, italic=False)


def _spacing(paragraph, before: int | None = None, after: int | None = None,
             line: float | None = None, align: Any = None, left: int | None = None,
             hanging: int | None = None, keep_next: bool | None = None) -> None:
    fmt = paragraph.paragraph_format
    if before is not None:
        fmt.space_before = Twips(before)
    if after is not None:
        fmt.space_after = Twips(after)
    if line is not None:
        fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    if left is not None:
        fmt.left_indent = Twips(left)
    if hanging is not None:
        fmt.first_line_indent = Twips(-hanging)
    if keep_next is not None:
        fmt.keep_with_next = keep_next


def _insert_ordered(parent, element, order: tuple[str, ...]) -> None:
    name = element.tag.split("}")[1]
    for existing in parent.findall(element.tag):
        parent.remove(existing)
    rank = order.index(name)
    for child in parent:
        child_name = child.tag.split("}")[1]
        if child_name in order and order.index(child_name) > rank:
            child.addprevious(element)
            return
    parent.append(element)


def _cell(cell, width: int, fill: str | None = None,
          margins: tuple[int, int, int, int] | None = None,
          borders: dict[str, tuple[int, str]] | None = None) -> None:
    cell.width = Twips(width)
    tc_pr = cell._tc.get_or_add_tcPr()
    if borders:
        element = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            if side in borders:
                size, colour = borders[side]
                edge = OxmlElement(f"w:{side}")
                edge.set(qn("w:val"), "single")
                edge.set(qn("w:sz"), str(size))
                edge.set(qn("w:space"), "0")
                edge.set(qn("w:color"), colour)
                element.append(edge)
        _insert_ordered(tc_pr, element, TCPR_ORDER)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        _insert_ordered(tc_pr, shd, TCPR_ORDER)
    if margins:
        element = OxmlElement("w:tcMar")
        for side, value in zip(("top", "left", "bottom", "right"), margins):
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:w"), str(value))
            edge.set(qn("w:type"), "dxa")
            element.append(edge)
        _insert_ordered(tc_pr, element, TCPR_ORDER)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _box(sz: int = 6) -> dict[str, tuple[int, str]]:
    return {side: (sz, BORDER) for side in ("top", "left", "bottom", "right")}


def _table(container, widths: list[int], rows: int = 1):
    if hasattr(container, "add_table") and container.__class__.__name__ in ("_Header", "_Footer"):
        table = container.add_table(rows=rows, cols=len(widths), width=Twips(sum(widths)))
    else:
        table = container.add_table(rows=rows, cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    for column, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        column.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
    return table


def _row_flags(row, header: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    if header:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _cell_paragraph(cell, first: bool):
    return cell.paragraphs[0] if first else cell.add_paragraph()


def _field(paragraph, instruction: str, size: float, colour: str) -> None:
    def fld(kind: str) -> None:
        run = _run(paragraph, "", size, colour)
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        run._r.append(element)

    fld("begin")
    run = _run(paragraph, "", size, colour)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    fld("separate")
    _run(paragraph, "1", size, colour)
    fld("end")


# --------------------------------------------------------------------------- #
# Document setup
# --------------------------------------------------------------------------- #
def _style_font(style, size: float, colour: str, bold: bool | None = None) -> None:
    style.font.name = "Arial"
    r_pr = style.element.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    fonts.set(qn("w:cs"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(colour)
    colour_el = r_pr.find(qn("w:color"))
    for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
        colour_el.attrib.pop(qn(attr), None)
    if bold is not None:
        style.font.bold = bold
    style.font.italic = False


def apply_house_styles(doc) -> None:
    styles = doc.styles
    specs = (
        ("Normal", 10, INK, None, None, 80, 1.08),
        ("Heading 1", 20, NAVY, True, 200, 100, None),
        ("Heading 2", 13, HEADING2, True, 160, 100, None),
        ("Heading 3", 11, BLUE, True, 160, 100, None),
        ("Title", 30, NAVY, True, 160, 100, 1.0),
        ("Caption", 8.5, MUTED, False, 60, 160, 1.0),
    )
    for name, size, colour, bold, before, after, line in specs:
        try:
            style = styles[name]
        except KeyError:
            continue
        _style_font(style, size, colour, bold)
        fmt = style.paragraph_format
        if before is not None:
            fmt.space_before = Twips(before)
        if after is not None:
            fmt.space_after = Twips(after)
        if line is not None:
            fmt.line_spacing = line
    title_ppr = styles["Title"].element.find(qn("w:pPr"))
    if title_ppr is not None:
        for border in title_ppr.findall(qn("w:pBdr")):
            title_ppr.remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Twips(11906), Twips(16838)
    section.top_margin, section.bottom_margin = Twips(935), Twips(907)
    section.left_margin, section.right_margin = Twips(935), Twips(935)
    section.header_distance = section.footer_distance = Twips(720)


def new_document(template: Path | None = None):
    if template is None:
        doc = Document()
        apply_house_styles(doc)
        return doc, False
    doc = Document(str(template))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return doc, True


def _has_content(part) -> bool:
    if part.is_linked_to_previous:
        return False
    return bool(part._element.xpath(".//w:t[normalize-space()] | .//w:drawing | .//w:pict"))


def write_header_footer(doc, customer: str, date_text: str, keep_existing: bool) -> None:
    section = doc.sections[0]
    if keep_existing and (_has_content(section.header) or _has_content(section.footer)):
        return
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = ""
    _spacing(paragraph, before=0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
    _run(paragraph, f"LOGICMONITOR  |  {customer.upper()} / INFOSYS HEALTH CHECK", 8, NAVY, True)
    footer = section.footer
    widths = [6803, 3231]
    table = _table(footer, widths)
    top_rule = {"top": (4, BORDER)}
    left, right = table.rows[0].cells
    _cell(left, widths[0], margins=(30, 0, 20, 0), borders=top_rule)
    _cell(right, widths[1], margins=(30, 0, 20, 0), borders=top_rule)
    _run(left.paragraphs[0], f"Internal | Assessment date: {date_text}", 8, MUTED, False)
    page = right.paragraphs[0]
    page.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(page, "Page ", 8, MUTED, False)
    _field(page, "PAGE", 8, MUTED)
    trailing = footer.paragraphs[0]
    footer._element.remove(trailing._p)
    footer._element.append(trailing._p)
    _spacing(trailing, before=0, after=0, line=1.0)


def set_properties(doc, assessment: Assessment) -> None:
    ctx = assessment.context
    props = doc.core_properties
    props.title = f"{ctx.customer} / Infosys - LM Envision Customer Health Check"
    props.subject = "Customer health check and Edwin AI readiness assessment"
    props.author = ctx.prepared_by or "Infosys"
    props.last_modified_by = ctx.prepared_by or "Infosys"
    props.keywords = "Infosys, LogicMonitor, LM Envision, Edwin AI, health check"
    props.comments = "Generated from the LogicMonitor REST API by the Infosys health check report tool."
    props.category = "Health check"
    stamp = ctx.generated_at.astimezone(timezone.utc).replace(tzinfo=None)
    props.created = props.modified = stamp
    props.revision = 1


# --------------------------------------------------------------------------- #
# Building blocks
# --------------------------------------------------------------------------- #
class Writer:
    def __init__(self, doc, template_mode: bool) -> None:
        self.doc = doc
        self.template_mode = template_mode

    def styled(self, style: str):
        try:
            return self.doc.add_paragraph(style=style)
        except KeyError:
            return self.doc.add_paragraph()

    def spacer(self, after: int = 0) -> None:
        paragraph = self.doc.add_paragraph()
        _spacing(paragraph, after=after)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def heading1(self, text: str, subtitle: str = "") -> None:
        paragraph = self.styled("Heading 1")
        _run(paragraph, text)
        if subtitle:
            sub = self.doc.add_paragraph()
            _spacing(sub, before=0, after=160, line=1.08, keep_next=False)
            _run(sub, subtitle, 9.5, MUTED, False)

    def heading2(self, text: str) -> None:
        _run(self.styled("Heading 2"), text)

    def body(self, text: str, ref: str = "") -> None:
        paragraph = self.doc.add_paragraph()
        _text(paragraph, text, 10, INK)
        if ref:
            _run(paragraph, f" {ref}", 7.5, BLUE, False, superscript=True)

    def note(self, text: str) -> None:
        paragraph = self.doc.add_paragraph()
        _text(paragraph, text, 9, MUTED)

    def numbered(self, items: list[str], size: float = 9.5) -> None:
        for index, item in enumerate(items, start=1):
            paragraph = self.doc.add_paragraph()
            _spacing(paragraph, before=0, after=60, line=1.08, left=408, hanging=295)
            _run(paragraph, f"{index}.", size, NAVY, True)
            _text(paragraph, f" {item}", size, INK)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            paragraph = self.styled("List Bullet")
            if paragraph.style.name != "List Bullet":
                _run(paragraph, "•  ", 9.5, NAVY, True)
            _spacing(paragraph, before=0, after=40, line=1.08, left=312, hanging=142)
            _text(paragraph, item, 9.5, INK)

    def banner(self, text: str) -> None:
        table = _table(self.doc, [10036])
        cell = table.rows[0].cells[0]
        _cell(cell, 10036, fill=NAVY, margins=(150, 180, 150, 180))
        _run(cell.paragraphs[0], text, 9, WHITE, True)

    def rag_callout(self, status: str, label: str, text: str) -> None:
        widths = [1871, 8050]
        table = _table(self.doc, widths)
        left, right = table.rows[0].cells
        fill, colour = status_colours(status)
        _cell(left, widths[0], fill=fill, margins=(170, 170, 170, 170), borders=_box())
        _cell(right, widths[1], fill=PALE, margins=(170, 170, 170, 170), borders=_box())
        paragraph = left.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(paragraph, status, 13, colour, True)
        paragraph = right.paragraphs[0]
        _run(paragraph, f"{label} ", 10.5, NAVY, True)
        _text(paragraph, text, 10.5, INK)

    def accent_callout(self, title: str, text: str, bar: str = BLUE, fill: str = PALE) -> None:
        widths = [142, 9780]
        table = _table(self.doc, widths)
        left, right = table.rows[0].cells
        _cell(left, widths[0], fill=bar, margins=(0, 0, 0, 0))
        _cell(right, widths[1], fill=fill, margins=(150, 180, 150, 180))
        heading = right.paragraphs[0]
        _spacing(heading, before=0, after=40, line=1.08)
        _run(heading, title, 10.5, NAVY, True)
        paragraph = right.add_paragraph()
        _spacing(paragraph, before=0, after=0, line=1.08)
        _text(paragraph, text, 10, INK)

    def fact(self, label: str, value: str) -> None:
        paragraph = self.doc.add_paragraph()
        _run(paragraph, label, 9, NAVY, True)
        _text(paragraph, f" {value}", 9, INK)

    def tiles(self, tiles: list[tuple[str, str, str]]) -> None:
        rows = (len(tiles) + 2) // 3
        table = _table(self.doc, [TILE_WIDTH] * 3, rows=rows)
        for index in range(rows * 3):
            cell = table.rows[index // 3].cells[index % 3]
            _cell(cell, TILE_WIDTH, fill=PALE, margins=(170, 145, 170, 145), borders=_box())
            if index >= len(tiles):
                continue
            value, label, sub = tiles[index]
            for position, (text, size, colour, bold) in enumerate(
                ((value, 18, NAVY, True), (label, 8.5, BLUE, True), (sub, 7.5, MUTED, False))
            ):
                paragraph = _cell_paragraph(cell, position == 0)
                _spacing(paragraph, before=0, after=20, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
                _text(paragraph, text, size, colour, bold)

    def cards(self, cards: list[tuple[str, str, list[str]]], title_size: float = 11.5,
              text_size: float = 8.5) -> None:
        table = _table(self.doc, [TILE_WIDTH] * 3)
        for cell, (title, colour, items) in zip(table.rows[0].cells, cards):
            _cell(cell, TILE_WIDTH, fill=PALE, margins=(0, 145, 160, 145), borders=_box())
            heading = cell.paragraphs[0]
            _spacing(heading, before=120, after=160, line=1.08, align=WD_ALIGN_PARAGRAPH.CENTER)
            _run(heading, title, title_size, colour, True)
            for item in items:
                paragraph = cell.add_paragraph()
                _spacing(paragraph, before=0, after=100, line=1.05, left=198, hanging=113)
                _run(paragraph, "• ", text_size + 0.5, colour, True)
                _text(paragraph, item, text_size, INK)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    def data_table(self, headers: list[str], rows: list[list[str]], widths: list[int],
                   aligns: list[str] | None = None, first_col: tuple[float, str] = (9, NAVY),
                   status_col: int | None = None, status_size: float = 8,
                   text_size: float = 8.5, header_size: float = 8.5,
                   first_col_fills: list[str] | None = None, last_col_colour: str = INK,
                   cell_margins: tuple[int, int, int, int] = (110, 120, 110, 120)) -> None:
        aligns = aligns or ["left"] * len(headers)
        align_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                     "right": WD_ALIGN_PARAGRAPH.RIGHT}
        table = _table(self.doc, widths, rows=len(rows) + 1)
        header = table.rows[0]
        _row_flags(header, header=True)
        for cell, text, width, align in zip(header.cells, headers, widths, aligns):
            _cell(cell, width, fill=NAVY, margins=cell_margins)
            paragraph = cell.paragraphs[0]
            _spacing(paragraph, before=0, after=0, line=1.05, align=align_map[align])
            _run(paragraph, text, header_size, WHITE, True)
        edges = {side: (4, BORDER) for side in ("left", "bottom", "right")}
        for row_index, values in enumerate(rows, start=1):
            row = table.rows[row_index]
            _row_flags(row)
            for col, (cell, value, width, align) in enumerate(zip(row.cells, values, widths, aligns)):
                fill = None
                colour, size, bold = INK, text_size, False
                if col == 0:
                    size, colour = first_col
                    bold = True
                    if first_col_fills:
                        fill = first_col_fills[(row_index - 1) % len(first_col_fills)]
                elif col == len(values) - 1:
                    colour = last_col_colour
                if status_col is not None and col == status_col:
                    fill, colour = status_colours(value)
                    size, bold = status_size, True
                _cell(cell, width, fill=fill, margins=cell_margins, borders=edges)
                paragraph = cell.paragraphs[0]
                _spacing(paragraph, before=0, after=0, line=1.05, align=align_map[align])
                _text(paragraph, str(value), size, colour, bold)

    def references(self, sources: list[tuple[str, str, str]]) -> None:
        widths = [567, 9354]
        table = _table(self.doc, widths, rows=len(sources))
        rule = {"bottom": (4, BORDER)}
        for row, (ref, title, desc) in zip(table.rows, sources):
            left, right = row.cells
            _cell(left, widths[0], fill=PALE_BLUE, margins=(90, 100, 90, 100), borders=rule)
            _cell(right, widths[1], margins=(90, 100, 90, 100), borders=rule)
            paragraph = left.paragraphs[0]
            _spacing(paragraph, before=0, after=0, line=1.05, align=WD_ALIGN_PARAGRAPH.CENTER)
            _run(paragraph, ref, 9, BLUE, True)
            heading = right.paragraphs[0]
            _spacing(heading, before=0, after=40, line=1.03)
            _text(heading, title, 9, NAVY, True)
            paragraph = right.add_paragraph()
            _spacing(paragraph, before=0, after=40, line=1.03)
            _text(paragraph, desc, 8.5, INK)

    def figure(self, png: bytes | None, caption: str, missing: str) -> None:
        if png:
            import io

            paragraph = self.doc.add_paragraph()
            _spacing(paragraph, before=0, after=20, line=1.08, align=WD_ALIGN_PARAGRAPH.CENTER)
            paragraph.add_run().add_picture(io.BytesIO(png), width=FIGURE_WIDTH)
        else:
            self.accent_callout("Figure not available", missing, bar=MUTED, fill=PALE_GREY)
        paragraph = self.styled("Caption")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(paragraph, caption, 8.5, MUTED, False)


# --------------------------------------------------------------------------- #
# Sections
# --------------------------------------------------------------------------- #
def _cover(w: Writer, a: Assessment) -> None:
    ctx = a.context
    w.banner("LOGICMONITOR  |  CUSTOMER HEALTH CHECK")
    w.spacer()
    paragraph = w.doc.add_paragraph()
    _spacing(paragraph, before=440, after=80, line=1.08)
    _run(paragraph, f"{ctx.customer.upper()} / INFOSYS", 14, BLUE, True)
    title = w.styled("Title")
    _spacing(title, before=0, after=80)
    _run(title, "LM Envision Customer Health Check", 30 if title.style.name != "Title" else None,
         NAVY if title.style.name != "Title" else None, True if title.style.name != "Title" else None)
    paragraph = w.doc.add_paragraph()
    _spacing(paragraph, before=0, after=440, line=1.08)
    _run(paragraph, ctx.subtitle, 12, MUTED, False)
    w.rag_callout(a.overall_status, "Overall condition:", a.overall_condition)
    w.spacer()
    w.spacer(after=40)
    w.accent_callout(a.conclusion_title, a.conclusion_text)
    w.spacer()
    w.spacer(after=80)
    for label, value in a.cover_facts:
        w.fact(label, value)


def _executive_summary(w: Writer, a: Assessment) -> None:
    w.heading1("Executive summary", "Current-state assessment, strategic context and immediate direction")
    for text, ref in a.summary_paragraphs:
        w.body(text, ref)
    w.heading2("Health at a glance")
    w.tiles(a.kpi_tiles)
    w.heading2("Immediate direction")
    w.numbered(a.immediate_direction)


def _scorecard(w: Writer, a: Assessment) -> None:
    w.heading1("Portal adoption and health scorecard", "Real-time portal metrics may differ from billed usage")
    w.data_table(
        ["Area", "Status", "Evidence and implication"],
        [[r.area, r.status, r.evidence] for r in a.scorecard],
        [1956, 1134, 6831],
        aligns=["left", "center", "left"],
        status_col=1,
        text_size=8.5,
    )
    w.spacer(after=80)
    w.accent_callout("Scorecard interpretation", a.interpretation)


def _good_weak_missing(w: Writer, a: Assessment) -> None:
    w.heading1(
        "What is good, what is weak and what is missing",
        "Assessment focused on business value, platform stickiness and expansion readiness",
    )
    w.cards([
        ("WHAT IS GOOD", GREEN_TEXT, a.good),
        ("WHAT IS WEAK", RED_TEXT, a.weak),
        ("WHAT IS MISSING", BLUE, a.missing),
    ])
    w.heading2("Stickiness and account economics")
    w.cards([
        ("Existing anchors", GREEN_TEXT, a.anchors),
        ("Expansion layers", BLUE, a.expansion),
        ("Retention risks", RED_TEXT, a.risks),
    ], title_size=10.5, text_size=8)


def _modules(w: Writer, a: Assessment) -> None:
    w.heading1("Module and capability usage", "Observed adoption versus value-realisation opportunity")
    w.data_table(
        ["Capability", "Observed usage", "Health", "Value motion"],
        [[r.capability, r.observed, r.health, r.motion] for r in a.modules],
        [1843, 2721, 1162, 4195],
        aligns=["left", "left", "center", "left"],
        status_col=2,
        status_size=7.5,
        text_size=8.5,
    )


def _edwin(w: Writer, a: Assessment) -> None:
    w.heading1("Edwin AI POV recommendation", a.edwin_subtitle)
    label, _, text = a.edwin_text.partition(": ")
    w.rag_callout(a.edwin_status, f"{label}:", text)
    w.heading2("Why the opportunity is credible")
    w.bullets(a.credible)
    w.heading2("Entry criteria that must be completed")
    w.bullets(a.entry_criteria)
    w.heading2("Recommended use cases and evidence")
    w.data_table(
        ["Use case", "Purpose", "Evidence to capture"],
        [list(row) for row in a.use_cases],
        [2608, 4082, 3231],
        first_col=(8.5, NAVY),
        header_size=8,
        text_size=8.5,
    )
    w.page_break()
    w.heading2("Three-week execution model")
    w.data_table(
        ["Horizon", "Priority actions", "Exit evidence"],
        [list(row) for row in a.execution],
        [1587, 5386, 2948],
        aligns=["center", "left", "left"],
        first_col=(10, NAVY),
        first_col_fills=[PALE_BLUE, PALE, PALE_GREY],
        last_col_colour=NAVY,
        header_size=8,
        cell_margins=(120, 120, 120, 120),
    )
    w.heading2("Core success measures")
    w.tiles(a.success_tiles)


def _plan(w: Writer, a: Assessment) -> None:
    w.heading1("30 / 60 / 90-day action plan", "Sequence foundational optimisation before broader AIOps scale")
    w.data_table(
        ["Horizon", "Priority actions", "Exit evidence"],
        [list(row) for row in a.plan],
        [1587, 5386, 2948],
        aligns=["center", "left", "left"],
        first_col=(10, NAVY),
        first_col_fills=[PALE_BLUE, PALE, PALE_GREY],
        last_col_colour=NAVY,
        header_size=8,
        cell_margins=(120, 120, 120, 120),
    )
    w.heading2("Decisions required at the onsite")
    w.numbered(a.decisions, size=9)
    w.spacer(after=80)
    w.accent_callout("Recommended customer positioning", a.positioning, bar=AMBER_TEXT, fill=AMBER_FILL)


def _sources(w: Writer, a: Assessment) -> None:
    w.heading1("Sources, assumptions and validation points", "Evidence used for this internal health check")
    w.references(a.sources)
    w.heading2("Validation points before customer use")
    w.bullets(a.validation)


def _appendix_a(w: Writer, a: Assessment, charts: dict[str, bytes | None], charts_note: str) -> None:
    m = a.metrics
    w.heading1("Appendix A - Portal evidence", "Selected charts supporting the key usage and risk observations")
    resources = m.get("resources")
    w.figure(
        charts.get("resources"),
        f"Figure A1. Resource footprint: {resources:,} monitored resources by type and hygiene state."
        if resources is not None else "Figure A1. Resource footprint.",
        charts_note or "Data unavailable: device inventory could not be read.",
    )
    prefix = "at least " if m.get("alerts_24h_partial") else ""
    w.figure(
        charts.get("alerts"),
        f"Figure A2. Alert volume: {prefix}{m['alerts_24h']:,} alerts raised in the 24 hours before collection, by severity."
        if m.get("alerts_24h") is not None else "Figure A2. Alert volume in the prior 24 hours.",
        charts_note or "Data unavailable: prior-day alerts could not be read.",
    )
    w.page_break()
    websites = m.get("websites")
    w.figure(
        charts.get("websites"),
        f"Figure A3. Website and synthetic checks: {websites:,} checks, {m.get('websites_external', 0):,} from external locations."
        if websites else "Figure A3. Website and synthetic checks.",
        charts_note or (
            "No website checks are configured." if websites == 0
            else "Data unavailable: websites could not be read."
        ),
    )
    collectors = m.get("collectors")
    w.figure(
        charts.get("collectors"),
        f"Figure A4. Collector estate: {collectors:,} collectors, {m.get('collectors_down', 0):,} down and "
        f"{m.get('collectors_no_failover', 0):,} carrying resources without failover."
        if collectors is not None else "Figure A4. Collector estate.",
        charts_note or "Data unavailable: collectors could not be read.",
    )
    w.note(
        "LM Logs and LM Uptime usage charts from the portal Usage page are not available through the "
        "REST API. " + "[Consultant to complete: add Usage page screenshots if required.]"
    )


def _appendix_b(w: Writer, a: Assessment) -> None:
    w.heading1("Appendix B - Supporting data", "Detail behind the alerting and collector findings")
    w.heading2("Top alert sources in the prior 24 hours")
    if a.top_sources:
        w.data_table(
            ["DataSource / alert type", "Alerts", "Share", "Most affected resource"],
            [[name, f"{count:,}", f"{share:.1f}%", device] for name, count, share, device in a.top_sources],
            [3600, 1100, 1100, 4121],
            aligns=["left", "right", "right", "left"],
            first_col=(8.5, NAVY),
            header_size=8,
        )
    else:
        w.note("No prior-day alerts were returned, or the alert data was unavailable.")
    w.heading2("Collector exceptions")
    exceptions = [r for r in a.collector_rows if r["issues"]]
    if exceptions:
        shown = exceptions[:25]
        w.data_table(
            ["Collector", "Status", "Size", "Resources", "Instances", "Findings"],
            [
                [
                    r["name"],
                    "RED" if r["down"] else "AMBER",
                    r["size"],
                    f"{r['hosts']:,}",
                    f"{r['instances']:,}",
                    "; ".join(r["issues"]),
                ]
                for r in shown
            ],
            [2300, 900, 1000, 1000, 1100, 3621],
            aligns=["left", "center", "left", "right", "right", "left"],
            first_col=(8.5, NAVY),
            status_col=1,
            header_size=8,
            text_size=8,
        )
        if len(exceptions) > len(shown):
            w.note(f"{len(exceptions) - len(shown)} further collectors have findings; see the operations health check.")
    elif a.collector_rows:
        w.note("No collector exceptions against the operations standards.")
    else:
        w.note("Data unavailable: collectors could not be read.")


def _appendix_c(w: Writer, a: Assessment) -> None:
    w.heading1("Appendix C - Data collection log", "Read-only REST API v3 endpoints queried for this report")
    w.data_table(
        ["Dataset", "Endpoint (GET)", "Records", "Status"],
        [list(row) for row in a.collection_log],
        [2200, 2900, 1000, 3821],
        aligns=["left", "left", "right", "left"],
        first_col=(8.5, NAVY),
        header_size=8,
        text_size=8,
        cell_margins=(60, 100, 60, 100),
    )


SECTIONS: list[tuple[str, Callable[..., None], bool]] = [
    ("Cover page", _cover, True),
    ("Executive summary", _executive_summary, True),
    ("Portal adoption and health scorecard", _scorecard, True),
    ("What is good, weak and missing", _good_weak_missing, True),
    ("Module and capability usage", _modules, True),
    ("Edwin AI POV recommendation", _edwin, True),
    ("30 / 60 / 90-day action plan", _plan, True),
    ("Sources and validation points", _sources, True),
    ("Appendix A - Portal evidence", _appendix_a, True),
    ("Appendix B - Supporting data", _appendix_b, True),
    ("Appendix C - Data collection log", _appendix_c, False),
]


def build_document(
    assessment: Assessment,
    charts: dict[str, bytes | None],
    template: Path | None = None,
    charts_note: str = "",
    progress: SectionProgress | None = None,
):
    doc, template_mode = new_document(template)
    writer = Writer(doc, template_mode)
    write_header_footer(doc, assessment.context.customer, assessment.context.date_text, template_mode)
    set_properties(doc, assessment)
    for name, render, page_break in SECTIONS:
        ok, note = True, ""
        try:
            if render is _appendix_a:
                render(writer, assessment, charts, charts_note)
            else:
                render(writer, assessment)
        except Exception as exc:  # one section must never stop the document
            ok, note = False, f"{exc.__class__.__name__}: {exc}"
            writer.body(f"Data unavailable: this section could not be generated ({note[:160]})")
        if progress:
            progress(name, ok, note)
        if page_break:
            writer.page_break()
    return doc
