#!/usr/bin/env python3
from __future__ import annotations

import io
import json
import os
import re
import sys
import tempfile
import unittest
import urllib.error
import urllib.parse
from contextlib import redirect_stderr, redirect_stdout
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

HERE = Path(__file__).resolve().parents[1]
OPS_DIR = HERE.parent / "Operations Portal Maintenance"
for _path in (HERE, OPS_DIR):
    if str(_path) not in sys.path:
        sys.path.insert(0, str(_path))

import lm_api  # noqa: E402
import portal_data  # noqa: E402
from assessment import (  # noqa: E402
    IST,
    ReportContext,
    build_assessment,
    count_of,
    load_rules,
)
from lm_api import ApiError, LmClient, LmConfig  # noqa: E402
from portal_data import (  # noqa: E402
    Dataset,
    FixtureSource,
    LiveSource,
    PortalData,
    collect_portal_data,
    endpoint_specs,
)

import lm_health_check_report as cli  # noqa: E402

try:
    import docx  # noqa: F401

    import charts
    import docx_writer

    HAVE_DOCX = True
except ImportError:  # pragma: no cover - python-docx is in requirements.txt
    HAVE_DOCX = False

FIXTURES = HERE / "tests" / "fixtures" / "api"
FAKE_KEY = "fake-access-key-must-never-be-printed"
FAKE_ENV = {
    "PATH": os.environ.get("PATH", "/usr/bin"),
    "LM_ACCOUNT": "example",
    "LM_ACCESS_ID": "fake-access-id",
    "LM_ACCESS_KEY": FAKE_KEY,
    "LM_CUSTOMER": "Example Corp",
}

# Headings of the Infosys health check layout, in order, followed by the
# appendices this generator adds.
EXPECTED_HEADINGS = [
    ("Heading 1", "Executive summary"),
    ("Heading 2", "Health at a glance"),
    ("Heading 2", "Immediate direction"),
    ("Heading 1", "Portal adoption and health scorecard"),
    ("Heading 1", "What is good, what is weak and what is missing"),
    ("Heading 2", "Stickiness and account economics"),
    ("Heading 1", "Module and capability usage"),
    ("Heading 1", "Edwin AI POV recommendation"),
    ("Heading 2", "Why the opportunity is credible"),
    ("Heading 2", "Entry criteria that must be completed"),
    ("Heading 2", "Recommended use cases and evidence"),
    ("Heading 2", "Three-week execution model"),
    ("Heading 2", "Core success measures"),
    ("Heading 1", "30 / 60 / 90-day action plan"),
    ("Heading 2", "Decisions required at the onsite"),
    ("Heading 1", "Sources, assumptions and validation points"),
    ("Heading 2", "Validation points before customer use"),
    ("Heading 1", "Appendix A - Portal evidence"),
    ("Heading 1", "Appendix B - Supporting data"),
    ("Heading 2", "Top alert sources in the prior 24 hours"),
    ("Heading 2", "Collector exceptions"),
    ("Heading 1", "Appendix C - Data collection log"),
]

EXPECTED_TABLE_HEADERS = [
    ["Area", "Status", "Evidence and implication"],
    ["Capability", "Observed usage", "Health", "Value motion"],
    ["Use case", "Purpose", "Evidence to capture"],
    ["Horizon", "Priority actions", "Exit evidence"],
    ["Horizon", "Priority actions", "Exit evidence"],
    ["DataSource / alert type", "Alerts", "Share", "Most affected resource"],
    ["Collector", "Status", "Size", "Resources", "Instances", "Findings"],
    ["Dataset", "Endpoint (GET)", "Records", "Status"],
]


# --------------------------------------------------------------------------- #
# Fake LogicMonitor portal (replaces urllib.request.urlopen)
# --------------------------------------------------------------------------- #
class FakeResponse:
    def __init__(self, payload: object) -> None:
        self._raw = json.dumps(payload).encode("utf-8")

    def read(self) -> bytes:
        return self._raw

    def __enter__(self) -> FakeResponse:
        return self

    def __exit__(self, *exc: object) -> bool:
        return False


def http_error(url: str, status: int, headers: dict | None = None) -> urllib.error.HTTPError:
    body = io.BytesIO(json.dumps({"errorMessage": f"synthetic HTTP {status}"}).encode("utf-8"))
    return urllib.error.HTTPError(url, status, "error", headers or {}, body)


class FakePortal:
    """Serves the fixtures the way REST API v3 does and records every request."""

    def __init__(self, failures: dict | None = None) -> None:
        self.failures = failures or {}
        self.requests: list[tuple[str, str, dict[str, str]]] = []
        self.routes: dict[str, str] = {}
        for spec in endpoint_specs(0, load_rules()):
            self.routes.setdefault(spec.path, spec.name)
        self.usage = json.loads((FIXTURES / "datasource_devices.json").read_text(encoding="utf-8"))

    def __call__(self, request, timeout=None) -> FakeResponse:
        url = urllib.parse.urlsplit(request.full_url)
        path = url.path.split("/santaba/rest", 1)[1]
        query = dict(urllib.parse.parse_qsl(url.query))
        self.requests.append((request.get_method(), path, query))
        failure = self.failures.get(path)
        status = failure(query) if callable(failure) else failure
        if status:
            raise http_error(request.full_url, status)
        usage = re.fullmatch(r"/setting/datasources/(\d+)/devices", path)
        if usage:
            return FakeResponse({"total": int(self.usage.get(usage.group(1), 0)), "items": []})
        name = self.routes.get(path)
        if path == "/alert/alerts":
            name = "alerts_24h" if 'cleared:"*"' in query.get("filter", "") else "alerts_active"
        if name is None:
            raise http_error(request.full_url, 404)
        payload = json.loads((FIXTURES / f"{name}.json").read_text(encoding="utf-8"))
        items, total = lm_api._items(payload), lm_api._total(payload)
        offset, size = int(query.get("offset", 0)), int(query.get("size", 50))
        body = {"total": len(items) if total is None else total, "items": items[offset : offset + size]}
        if "data" in payload:
            body = {"status": 200, "errmsg": "OK", "data": body}
        return FakeResponse(body)

    def paths(self) -> list[str]:
        return [path for _, path, _ in self.requests]


def fake_client(max_retries: int = 4) -> LmClient:
    return LmClient(LmConfig("example", "fake-access-id", FAKE_KEY), max_retries=max_retries)


def fixture_data() -> tuple[PortalData, float]:
    source = FixtureSource(FIXTURES)
    now = float(source.meta["now"])
    return collect_portal_data(source, now, load_rules(), clone_pattern=_clone_pattern()), now


def _clone_pattern():
    return re.compile(load_rules()["clone_name_pattern"]).search


def context(now: float) -> ReportContext:
    return ReportContext(
        customer="Example Corp",
        portal="example.logicmonitor.com",
        generated_at=datetime.fromtimestamp(now, IST),
        prepared_by="Test Consultant",
    )


def failing_data(now: float, reason: str = "HTTP 403: the API token's role does not have view permission") -> PortalData:
    datasets = {
        spec.name: Dataset(spec.name, spec.label, spec.path, error=reason)
        for spec in endpoint_specs(now, load_rules())
    }
    return PortalData(portal="example.logicmonitor.com", collected_at=now, datasets=datasets)


# --------------------------------------------------------------------------- #
# Shared client
# --------------------------------------------------------------------------- #
class ClientTests(unittest.TestCase):
    def test_collect_follows_negative_total(self) -> None:
        pages = [[{"id": 1}, {"id": 2}], [{"id": 3}, {"id": 4}], [{"id": 5}]]
        calls: list[dict[str, str]] = []

        class PagedClient(LmClient):
            def get(self, resource_path, query=None):
                calls.append(dict(query or {}))
                index = int(query["offset"]) // 2
                return {"total": -3, "items": pages[index] if index < len(pages) else []}

        with patch.object(lm_api, "PAGE_SIZE", 2):
            items = PagedClient(LmConfig("example", "id", "key")).collect("/alert/alerts")
        self.assertEqual([i["id"] for i in items], [1, 2, 3, 4, 5])
        self.assertEqual(len(calls), 3)

    def test_429_honours_retry_after(self) -> None:
        responses = [http_error("https://x", 429, {"Retry-After": "3"}), FakeResponse({"total": 0, "items": []})]

        def urlopen(request, timeout=None):
            result = responses.pop(0)
            if isinstance(result, Exception):
                raise result
            return result

        with patch.object(lm_api.urllib.request, "urlopen", urlopen), patch.object(lm_api.time, "sleep") as sleep:
            payload = fake_client().get("/device/devices", {"size": "1"})
        self.assertEqual(payload, {"total": 0, "items": []})
        sleep.assert_called_once_with(3.0)

    def test_backoff_is_capped(self) -> None:
        exc = http_error("https://x", 429, {"Retry-After": "3600"})
        self.assertEqual(lm_api._backoff_seconds(exc, 1), lm_api.MAX_BACKOFF_SECONDS)
        self.assertEqual(lm_api._backoff_seconds(http_error("https://x", 503), 2), 4.0)

    def test_write_verbs_are_refused(self) -> None:
        with self.assertRaises(ApiError):
            fake_client()._request("DELETE", "/device/devices/1")


# --------------------------------------------------------------------------- #
# Collection
# --------------------------------------------------------------------------- #
class CollectionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.rules = load_rules()
        self.now = float(FixtureSource(FIXTURES).meta["now"])

    def collect(self, portal: FakePortal) -> PortalData:
        with patch.object(lm_api.urllib.request, "urlopen", portal), patch.object(lm_api.time, "sleep"):
            return collect_portal_data(LiveSource(fake_client()), self.now, self.rules, clone_pattern=_clone_pattern())

    def test_fixtures_cover_every_endpoint_in_both_shapes(self) -> None:
        shapes = set()
        for spec in endpoint_specs(self.now, self.rules):
            path = FIXTURES / f"{spec.name}.json"
            self.assertTrue(path.is_file(), spec.name)
            shapes.add("data" in json.loads(path.read_text(encoding="utf-8")))
        self.assertEqual(shapes, {True, False})

    def test_live_source_matches_fixture_source(self) -> None:
        portal = FakePortal()
        live = self.collect(portal)
        expected, _ = fixture_data()
        self.assertEqual(len(live.datasets), 28)
        for name, dataset in expected.datasets.items():
            self.assertTrue(live.datasets[name].ok, f"{name}: {live.datasets[name].error}")
            self.assertEqual(live.datasets[name].count, dataset.count, name)
        self.assertEqual(live.module_usage, expected.module_usage)
        self.assertEqual({verb for verb, _, _ in portal.requests}, {"GET"})

    def test_alert_windows_and_pagination(self) -> None:
        portal = FakePortal()
        with patch.object(portal_data, "PAGE_SIZE", 100):
            data = self.collect(portal)
        self.assertEqual(data.get("alerts_24h").count, 372)
        self.assertEqual(data.get("alerts_active").count, 130)
        windows = [q for _, p, q in portal.requests if p == "/alert/alerts" and 'cleared:"*"' in q.get("filter", "")]
        self.assertEqual([q["offset"] for q in windows], ["0", "100", "200", "300"])
        self.assertIn("startEpoch>:", windows[0]["filter"])

    def test_forbidden_endpoint_is_recorded_not_fatal(self) -> None:
        data = self.collect(FakePortal({"/setting/admins": 403}))
        users = data.get("users")
        self.assertFalse(users.ok)
        self.assertIn("HTTP 403", users.error)
        self.assertTrue(data.get("api_tokens").ok)
        self.assertEqual(sum(1 for d in data.datasets.values() if not d.ok), 1)

    def test_rejected_fields_are_retried_without_fields(self) -> None:
        portal = FakePortal({"/setting/integrations": lambda q: 400 if "fields" in q else None})
        data = self.collect(portal)
        self.assertTrue(data.get("integrations").ok)
        queries = [q for _, p, q in portal.requests if p == "/setting/integrations"]
        self.assertEqual(len(queries), 2)
        self.assertIn("fields", queries[0])
        self.assertNotIn("fields", queries[1])

    def test_error_mid_paging_keeps_partial_data(self) -> None:
        def fail_third_page(query):
            return 403 if 'cleared:"*"' in query.get("filter", "") and query.get("offset") == "200" else None

        with patch.object(portal_data, "PAGE_SIZE", 100):
            data = self.collect(FakePortal({"/alert/alerts": fail_third_page}))
        day = data.get("alerts_24h")
        self.assertTrue(day.ok)
        self.assertEqual(len(day.items), 200)
        self.assertIn("stopped after 200", day.partial)


# --------------------------------------------------------------------------- #
# Assessment rules
# --------------------------------------------------------------------------- #
class AssessmentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.data, cls.now = fixture_data()
        cls.assessment = build_assessment(cls.data, context(cls.now))
        cls.rows = {r.key: r for r in cls.assessment.scorecard}

    def test_scorecard_statuses(self) -> None:
        self.assertEqual(len(self.assessment.scorecard), 18)
        for row in self.assessment.scorecard:
            self.assertNotIn("could not evaluate", row.evidence, row.key)
        expected = {
            "alerting": "RED",
            "collectors": "AMBER / RED",
            "footprint": "AMBER",
            "routing": "AMBER",
            "sdt": "AMBER",
            "access": "AMBER",
            "discovery": "AMBER",
            "telemetry": "GREEN",
            "datapoints": "UNKNOWN",
            "logs": "UNKNOWN",
        }
        for key, status in expected.items():
            self.assertEqual(self.rows[key].status, status, key)
        self.assertEqual(self.assessment.overall_status, "AMBER")

    def test_metrics_from_fixtures(self) -> None:
        m = self.assessment.metrics
        self.assertEqual(m["resources"], 147)
        self.assertEqual(m["services"], 6)
        self.assertEqual(m["collectors_down"], 1)
        self.assertEqual(m["alerts_24h"], 372)
        self.assertEqual(m["clones_checked"], 4)
        self.assertEqual(m["clones_unused"], 2)
        self.assertEqual(self.assessment.top_sources[0][0], "SNMP_Network_Interfaces")
        self.assertNotIn("BGP-", [name for name, *_ in self.assessment.top_sources])

    def test_unavailable_api_areas_are_placeholders(self) -> None:
        self.assertIn("[Consultant to complete:", self.rows["core"].evidence)
        self.assertIn("[Consultant to complete:", self.rows["datapoints"].evidence)
        self.assertTrue(any("[Consultant to complete:" in text for text, _ in self.assessment.summary_paragraphs))

    def test_contracted_resources_enable_headroom(self) -> None:
        ctx = context(self.now)
        ctx.contracted_resources = 200
        rows = {r.key: r for r in build_assessment(self.data, ctx).scorecard}
        self.assertEqual(rows["core"].status, "GREEN")
        self.assertIn("53 resources of headroom", rows["core"].evidence)

    def test_every_endpoint_failing_degrades_gracefully(self) -> None:
        assessment = build_assessment(failing_data(self.now), context(self.now))
        rows = {r.key: r for r in assessment.scorecard}
        self.assertEqual(rows["footprint"].status, "UNKNOWN")
        self.assertTrue(rows["footprint"].evidence.startswith("Data unavailable: Devices: HTTP 403"))
        self.assertTrue(all(r.status == "UNKNOWN" for r in assessment.scorecard), [r.status for r in assessment.scorecard])
        for row in assessment.scorecard:
            self.assertNotIn("could not evaluate", row.evidence, row.key)
        self.assertEqual(assessment.overall_status, "UNKNOWN")

    def test_singular_wording(self) -> None:
        self.assertEqual(count_of(1, "collector"), "1 collector")
        self.assertEqual(count_of(2, "priority", "priorities"), "2 priorities")
        summary = " ".join(text for text, _ in self.assessment.summary_paragraphs)
        self.assertIn("1 of 10 collectors is down", summary)


# --------------------------------------------------------------------------- #
# Word document
# --------------------------------------------------------------------------- #
def headings(document) -> list[tuple[str, str]]:
    return [(p.style.name, p.text) for p in document.paragraphs if p.style.name.startswith("Heading")]


def table_headers(document) -> list[list[str]]:
    """First-row text of every data table (tables with a repeating header row)."""
    return [
        [c.text for c in table.rows[0].cells]
        for table in document.tables
        if table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    ]


def all_text(document) -> str:
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


@unittest.skipUnless(HAVE_DOCX, "python-docx is not installed")
class DocumentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.data, cls.now = fixture_data()
        cls.assessment = build_assessment(cls.data, context(cls.now))
        m = cls.assessment.metrics
        cls.charts = {
            "resources": charts.resource_chart(m),
            "alerts": charts.alert_chart(m),
            "websites": charts.website_chart(m),
            "collectors": charts.collector_chart(cls.assessment.collector_rows),
        }
        cls.sections: list[tuple[str, bool]] = []
        cls.document = docx_writer.build_document(
            cls.assessment, cls.charts, progress=lambda name, ok, note: cls.sections.append((name, ok))
        )

    def test_heading_order_matches_layout(self) -> None:
        self.assertEqual(headings(self.document), EXPECTED_HEADINGS)
        title = [p.text for p in self.document.paragraphs if p.style.name == "Title"]
        self.assertEqual(title, ["LM Envision Customer Health Check"])

    def test_tables_have_expected_columns(self) -> None:
        self.assertEqual(table_headers(self.document), EXPECTED_TABLE_HEADERS)
        scorecard = next(t for t in self.document.tables if t.rows[0].cells[0].text == "Area")
        self.assertEqual(len(scorecard.rows), 19)
        modules = next(t for t in self.document.tables if t.rows[0].cells[0].text == "Capability")
        self.assertEqual(len(modules.rows), 11)

    def test_cover_tiles_and_sources(self) -> None:
        text = all_text(self.document)
        self.assertIn("EXAMPLE CORP / INFOSYS", text)
        self.assertIn("Overall condition:", text)
        self.assertIn("Primary conclusion", text)
        for label in ("Devices", "Instances", "Alerts in prior day", "Open alerts", "Services", "Collectors"):
            self.assertIn(label, text)
        self.assertIn("LM Envision REST API v3 - example.logicmonitor.com", text)
        self.assertIn("[Consultant to complete:", text)

    @unittest.skipUnless(HAVE_DOCX and charts.charts_available(), "matplotlib is not installed")
    def test_appendix_has_four_charts(self) -> None:
        self.assertEqual(len(self.document.inline_shapes), 4)
        captions = [p.text for p in self.document.paragraphs if p.style.name == "Caption"]
        self.assertEqual([c.split(".")[0] for c in captions], ["Figure A1", "Figure A2", "Figure A3", "Figure A4"])

    def test_header_footer_and_page_setup(self) -> None:
        section = self.document.sections[0]
        self.assertTrue(section.different_first_page_header_footer)
        self.assertIn("EXAMPLE CORP / INFOSYS HEALTH CHECK", section.header.paragraphs[0].text)
        footer = " ".join(c.text for t in section.footer.tables for c in t.rows[0].cells)
        self.assertIn("Internal | Assessment date: 20 October 2026", footer)
        self.assertIn("Page", footer)
        self.assertEqual((section.page_width.twips, section.page_height.twips), (11906, 16838))
        self.assertEqual(self.document.styles["Normal"].font.name, "Arial")

    def test_every_section_reports_ok(self) -> None:
        self.assertEqual([name for name, _ in self.sections], [name for name, _, _ in docx_writer.SECTIONS])
        self.assertTrue(all(ok for _, ok in self.sections))

    def test_failing_portal_still_produces_a_document(self) -> None:
        assessment = build_assessment(failing_data(self.now), context(self.now))
        document = docx_writer.build_document(assessment, {})
        self.assertEqual(headings(document), EXPECTED_HEADINGS)
        self.assertIn("Data unavailable: Devices: HTTP 403", all_text(document))

    def test_section_error_is_contained(self) -> None:
        def broken(writer, assessment):
            raise ValueError("synthetic failure")

        sections = [(n, broken if n == "Module and capability usage" else f, b) for n, f, b in docx_writer.SECTIONS]
        notes: list[tuple[str, bool]] = []
        with patch.object(docx_writer, "SECTIONS", sections):
            document = docx_writer.build_document(
                self.assessment, self.charts, progress=lambda name, ok, note: notes.append((name, ok))
            )
        self.assertIn(("Module and capability usage", False), notes)
        self.assertIn("Data unavailable: this section could not be generated (ValueError: synthetic failure)", all_text(document))
        self.assertIn(("Heading 1", "Edwin AI POV recommendation"), headings(document))

    def test_template_keeps_its_header_and_drops_its_body(self) -> None:
        template = docx.Document()
        template.sections[0].header.paragraphs[0].text = "TEMPLATE HEADER"
        template.add_paragraph("TEMPLATE BODY TEXT")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "template.docx"
            template.save(str(path))
            document = docx_writer.build_document(self.assessment, self.charts, template=path)
        self.assertEqual(document.sections[0].header.paragraphs[0].text, "TEMPLATE HEADER")
        self.assertNotIn("TEMPLATE BODY TEXT", all_text(document))
        self.assertEqual(headings(document), EXPECTED_HEADINGS)


# --------------------------------------------------------------------------- #
# Command line
# --------------------------------------------------------------------------- #
def run_cli(argv: list[str], env: dict[str, str], urlopen=None) -> tuple[int, str, str]:
    stdout, stderr = io.StringIO(), io.StringIO()
    patches = [
        patch.dict("os.environ", env, clear=True),
        patch.object(cli, "load_dotenv", lambda *a, **k: None),
        patch.object(lm_api.time, "sleep"),
    ]
    if urlopen is not None:
        patches.append(patch.object(lm_api.urllib.request, "urlopen", urlopen))
    for item in patches:
        item.start()
    try:
        with redirect_stdout(stdout), redirect_stderr(stderr):
            code = cli.main(argv)
    finally:
        for item in reversed(patches):
            item.stop()
    return code, stdout.getvalue(), stderr.getvalue()


class CliTests(unittest.TestCase):
    def test_report_filename(self) -> None:
        moment = datetime(2026, 10, 20, 9, 0, tzinfo=IST)
        self.assertEqual(
            cli.report_filename("Example Corp", moment),
            "Example_Corp_Infosys_LM_Envision_Health_Check_20_Oct_2026.docx",
        )
        self.assertNotIn("/", cli.report_filename("A/B & Co.", moment))

    def test_parse_portal(self) -> None:
        self.assertEqual(cli.parse_portal("https://acme.logicmonitor.com/santaba/", "logicmonitor.com"), ("acme", "logicmonitor.com"))
        self.assertEqual(cli.parse_portal("acme", "lmgov.us"), ("acme", "lmgov.us"))
        with self.assertRaises(cli.InputError):
            cli.parse_portal("not a portal!", "logicmonitor.com")

    def test_missing_inputs_exit_two_without_secrets(self) -> None:
        env = {"PATH": "/usr/bin", "LM_ACCESS_KEY": FAKE_KEY}
        code, out, err = run_cli(["--no-prompt"], env)
        self.assertEqual(code, 2)
        self.assertIn("LM_ACCOUNT", err)
        self.assertIn("LM_ACCESS_ID", err)
        self.assertIn("LM_CUSTOMER", err)
        self.assertNotIn("LM_ACCESS_KEY or the prompt", err)
        self.assertNotIn(FAKE_KEY, out + err)
        self.assertNotIn("Traceback", err)

    def test_prompts_read_the_key_without_echo(self) -> None:
        args = cli.parse_args([])
        answers = iter(["example", "fake-access-id", "Example Corp", ""])
        with patch.dict("os.environ", {"PATH": "/usr/bin"}, clear=True), patch.object(
            cli, "_interactive", return_value=True
        ), patch("builtins.input", lambda prompt="": next(answers)), patch.object(
            cli.getpass, "getpass", return_value=FAKE_KEY
        ) as hidden:
            values = cli.resolve_inputs(args)
        hidden.assert_called_once()
        self.assertEqual(values["access_key"], FAKE_KEY)
        self.assertEqual((values["account"], values["domain"]), ("example", "logicmonitor.com"))

    def test_unreachable_portal_is_a_clean_error(self) -> None:
        def unreachable(request, timeout=None):
            raise urllib.error.URLError("[Errno -2] Name or service not known")

        env = dict(FAKE_ENV, LM_ACCOUNT="nonexistent-portal", LM_DOMAIN="invalid.example")
        code, out, err = run_cli(["--no-prompt"], env, unreachable)
        self.assertEqual(code, 1)
        self.assertIn("Connection failed: Could not reach https://nonexistent-portal.invalid.example", err)
        self.assertNotIn("Traceback", err)
        self.assertNotIn(FAKE_KEY, out + err)

    def test_bad_credentials_are_a_clean_error(self) -> None:
        code, out, err = run_cli(["--no-prompt"], FAKE_ENV, FakePortal({"/device/devices": 401}))
        self.assertEqual(code, 1)
        self.assertIn("Authentication failed (HTTP 401)", err)
        self.assertNotIn(FAKE_KEY, out + err)

    @unittest.skipUnless(HAVE_DOCX, "python-docx is not installed")
    def test_end_to_end_against_fake_portal(self) -> None:
        portal = FakePortal({"/setting/accesslogs": 403})
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "report.docx"
            code, out, err = run_cli(["--no-prompt", "--output", str(output)], FAKE_ENV, portal)
            self.assertEqual(code, 0, err)
            self.assertTrue(output.is_file())
            document = docx.Document(str(output))
        self.assertEqual(headings(document), EXPECTED_HEADINGS)
        self.assertIn("[28/28] Audit log", out)
        self.assertIn("unavailable - HTTP 403", out)
        self.assertIn("1 of 28 endpoints were unavailable", out)
        self.assertIn("Report written:", out)
        self.assertNotIn(FAKE_KEY, out + err)
        self.assertEqual({verb for verb, _, _ in portal.requests}, {"GET"})

    @unittest.skipUnless(HAVE_DOCX, "python-docx is not installed")
    def test_demo_mode(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "demo"
            code, out, err = run_cli(["--demo", "--output", str(output)], {"PATH": "/usr/bin"})
            self.assertEqual(code, 0, err)
            self.assertTrue(output.with_suffix(".docx").is_file())
        self.assertIn("Overall status: AMBER", out)

    def test_missing_template_is_an_input_error(self) -> None:
        code, _, err = run_cli(["--demo", "--template", "/nonexistent/template.docx"], {"PATH": "/usr/bin"})
        self.assertEqual(code, 2)
        self.assertIn("Template not found", err)


if __name__ == "__main__":
    unittest.main()
