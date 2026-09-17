#!/usr/bin/env python3
"""
lm_api.py
=========
Shared LogicMonitor REST API v3 client and helpers.

Imported by:
    lm_threshold_apply.py
    lm_alert_rules.py
    lm_escalation_chains.py

Not run directly. Keep this file in the same folder as those scripts.

Provides:
    LogicMonitorClient   LMv1-signed client with retry, paging and lookups
    LogicMonitorError    single exception type for API failures
    prompt_credentials   interactive portal / Access ID / Access Key prompt
    read_sheet           Excel sheet reader returning list[dict]
    configure_logging    console + file logging
    write_report         per-row outcome CSV
    RowResult            per-row outcome record
"""

from __future__ import annotations

import base64
import csv
import getpass
import hashlib
import hmac
import json
import logging
import os
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

try:
    import requests
except ImportError:
    sys.exit("Missing dependency. Install with:  pip install requests")

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Missing dependency. Install with:  pip install openpyxl")

API_VERSION = "3"
REQUEST_TIMEOUT = 45
MAX_RETRIES = 4
BACKOFF_BASE = 2.0
PAGE_SIZE = 1000

log = logging.getLogger("lm")


class LogicMonitorError(RuntimeError):
    """Any non-retryable API or input failure."""


# --------------------------------------------------------------------------- #
# Client
# --------------------------------------------------------------------------- #
class LogicMonitorClient:
    def __init__(self, company: str, access_id: str, access_key: str):
        self.company = company
        self.base_url = f"https://{company}.logicmonitor.com/santaba/rest"
        self._access_id = access_id
        self._access_key = access_key.encode("utf-8")
        self._session = requests.Session()
        self._cache: dict[str, Any] = {}

    # -- auth ------------------------------------------------------------- #
    def _auth_header(self, verb: str, resource_path: str, body: str) -> str:
        epoch = str(int(time.time() * 1000))
        message = f"{verb}{epoch}{body}{resource_path}"
        digest = hmac.new(self._access_key, message.encode("utf-8"), hashlib.sha256).hexdigest()
        signature = base64.b64encode(digest.encode("utf-8")).decode("utf-8")
        return f"LMv1 {self._access_id}:{signature}:{epoch}"

    # -- transport -------------------------------------------------------- #
    def request(
        self,
        verb: str,
        resource_path: str,
        payload: dict[str, Any] | None = None,
        params: dict[str, Any] | None = None,
    ) -> Any:
        """Single API call. Query params are excluded from the signature, per LM docs."""
        body = json.dumps(payload) if payload is not None else ""
        url = self.base_url + resource_path
        params = dict(params or {})
        params["v"] = API_VERSION

        for attempt in range(1, MAX_RETRIES + 1):
            headers = {
                "Content-Type": "application/json",
                "X-Version": API_VERSION,
                "Authorization": self._auth_header(verb, resource_path, body),
            }
            try:
                response = self._session.request(
                    verb,
                    url,
                    data=body if body else None,
                    params=params,
                    headers=headers,
                    timeout=REQUEST_TIMEOUT,
                )
            except requests.RequestException as exc:
                if attempt == MAX_RETRIES:
                    raise LogicMonitorError(f"Network error on {resource_path}: {exc}") from exc
                self._sleep(attempt, exc.__class__.__name__)
                continue

            if response.status_code == 429 or 500 <= response.status_code < 600:
                if attempt == MAX_RETRIES:
                    raise LogicMonitorError(
                        f"{verb} {resource_path} failed after {MAX_RETRIES} attempts "
                        f"(HTTP {response.status_code}): {response.text[:300]}"
                    )
                self._sleep(attempt, f"HTTP {response.status_code}")
                continue

            if response.status_code == 401:
                raise LogicMonitorError(
                    "Authentication failed (HTTP 401). Check the portal name, Access ID "
                    "and Access Key, and confirm the API token is enabled."
                )
            if response.status_code == 403:
                raise LogicMonitorError(
                    f"Permission denied (HTTP 403) on {resource_path}. The API token's role "
                    "is missing a required right."
                )

            try:
                parsed = response.json()
            except ValueError:
                raise LogicMonitorError(
                    f"{verb} {resource_path} returned non-JSON (HTTP {response.status_code})"
                ) from None

            if not response.ok or parsed.get("errmsg", "OK") != "OK":
                raise LogicMonitorError(
                    f"{verb} {resource_path} -> HTTP {response.status_code}: "
                    f"{parsed.get('errmsg') or parsed}"
                )
            return parsed.get("data", parsed)

        raise LogicMonitorError(f"{verb} {resource_path} exhausted retries")

    @staticmethod
    def _sleep(attempt: int, reason: str) -> None:
        wait = BACKOFF_BASE ** attempt
        log.warning("Retry %d/%d in %.1fs (%s)", attempt, MAX_RETRIES, wait, reason)
        time.sleep(wait)

    def get_items(self, resource_path: str, params: dict[str, Any] | None = None) -> list[dict]:
        """GET a collection, following paging until exhausted."""
        collected: list[dict] = []
        offset = 0
        while True:
            page = dict(params or {})
            page.update({"size": PAGE_SIZE, "offset": offset})
            data = self.request("GET", resource_path, params=page)
            items = data.get("items", []) if isinstance(data, dict) else []
            collected.extend(items)
            total = data.get("total", len(collected)) if isinstance(data, dict) else len(collected)
            offset += len(items)
            if not items or offset >= total or len(items) < PAGE_SIZE:
                return collected

    # -- lookups ---------------------------------------------------------- #
    def find_device(self, identifier: str) -> dict:
        """Resolve a device by displayName first, then by name. Raises if not unique."""
        key = f"device:{identifier}"
        if key in self._cache:
            return self._cache[key]

        for field_name in ("displayName", "name"):
            items = self.get_items(
                "/device/devices",
                params={
                    "filter": f'{field_name}:"{identifier}"',
                    "fields": "id,name,displayName,preferredCollectorId",
                },
            )
            if len(items) == 1:
                self._cache[key] = items[0]
                return items[0]
            if len(items) > 1:
                raise LogicMonitorError(
                    f"'{identifier}' matches {len(items)} devices by {field_name}. "
                    "Use the exact display name."
                )
        raise LogicMonitorError(f"Device not found: {identifier}")

    def find_device_datasource(self, device_id: int, datasource_name: str) -> dict:
        """Resolve the applied DataSource (hdsId) on a device by DataSource name."""
        key = f"dds:{device_id}:{datasource_name}"
        if key in self._cache:
            return self._cache[key]

        items = self.get_items(
            f"/device/devices/{device_id}/devicedatasources",
            params={"fields": "id,dataSourceName,dataSourceId,instanceNumber"},
        )
        wanted = datasource_name.strip().lower()
        matches = [i for i in items if (i.get("dataSourceName") or "").lower() == wanted]
        if not matches:
            available = ", ".join(sorted(i.get("dataSourceName", "") for i in items)[:12])
            raise LogicMonitorError(
                f"DataSource '{datasource_name}' is not applied to this device. "
                f"Applied DataSources include: {available or 'none'}"
            )
        self._cache[key] = matches[0]
        return matches[0]

    def list_instances(self, device_id: int, hds_id: int) -> list[dict]:
        key = f"inst:{device_id}:{hds_id}"
        if key not in self._cache:
            self._cache[key] = self.get_items(
                f"/device/devices/{device_id}/devicedatasources/{hds_id}/instances",
                params={"fields": "id,name,displayName,groupId,groupName,stopMonitoring"},
            )
        return self._cache[key]

    def find_escalation_chain(self, name: str) -> dict:
        key = f"chain:{name.lower()}"
        if key in self._cache:
            return self._cache[key]

        items = self.get_items("/setting/alert/chains", params={"fields": "id,name,description"})
        wanted = name.strip().lower()
        matches = [c for c in items if (c.get("name") or "").strip().lower() == wanted]
        if not matches:
            raise LogicMonitorError(f"Escalation chain not found: {name}")
        self._cache[key] = matches[0]
        return matches[0]

    def find_alert_rule(self, name: str) -> dict | None:
        items = self.get_items("/setting/alert/rules", params={"fields": "id,name,priority"})
        wanted = name.strip().lower()
        for rule in items:
            if (rule.get("name") or "").strip().lower() == wanted:
                return rule
        return None

    def find_device_group(self, full_path: str) -> dict:
        clean = full_path.strip().strip("/")
        key = f"group:{clean.lower()}"
        if key in self._cache:
            return self._cache[key]

        items = self.get_items(
            "/device/groups", params={"filter": f'fullPath:"{clean}"', "fields": "id,name,fullPath"}
        )
        if not items:
            raise LogicMonitorError(f"Resource group not found: {full_path}")
        self._cache[key] = items[0]
        return items[0]


# --------------------------------------------------------------------------- #
# Credentials
# --------------------------------------------------------------------------- #
def prompt_credentials(company: str | None = None) -> LogicMonitorClient:
    """Prompt for portal, Access ID and Access Key. Env vars skip the prompts."""
    company = company or os.getenv("LM_COMPANY") or input(
        "LogicMonitor portal / company name (e.g. infosys): "
    ).strip()
    access_id = os.getenv("LM_ACCESS_ID") or input("Access ID: ").strip()
    access_key = os.getenv("LM_ACCESS_KEY") or getpass.getpass("Access Key (hidden): ").strip()

    if not (company and access_id and access_key):
        raise LogicMonitorError("Portal name, Access ID and Access Key are all required.")
    log.info("Portal: %s.logicmonitor.com", company)
    return LogicMonitorClient(company, access_id, access_key)


# --------------------------------------------------------------------------- #
# Excel input
# --------------------------------------------------------------------------- #
def read_sheet(path: Path, sheet: str | None = None, required: tuple[str, ...] = ()) -> list[dict]:
    """
    Read the first (or named) worksheet into a list of dicts.

    Row 1 is the header. Values are trimmed strings.

    Reading stops at the first completely blank row, so anything written below
    the data block (notes, working columns) is never parsed. Rows whose first
    populated cell starts with '#' are treated as comments and skipped.
    """
    if not path.is_file():
        raise LogicMonitorError(f"Excel file not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]
    except KeyError:
        raise LogicMonitorError(
            f"Sheet '{sheet}' not found. Sheets present: {', '.join(workbook.sheetnames)}"
        ) from None

    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        raise LogicMonitorError(f"{path.name} is empty.") from None

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    missing = [column for column in required if column not in headers]
    if missing:
        raise LogicMonitorError(
            f"{path.name} sheet '{worksheet.title}' is missing required column(s): "
            f"{', '.join(missing)}. Columns found: {', '.join(c for c in headers if c)}"
        )

    records: list[dict] = []
    for number, values in enumerate(rows, start=2):
        record = {
            headers[i]: ("" if value is None else str(value).strip())
            for i, value in enumerate(values)
            if i < len(headers) and headers[i]
        }
        if not any(record.values()):
            break
        first = next((v for v in record.values() if v), "")
        if first.startswith("#"):
            continue
        record["_row"] = str(number)
        records.append(record)

    workbook.close()
    return records


def as_bool(value: str) -> bool:
    return value.strip().lower() in {"true", "yes", "y", "1", "enable", "enabled", "on"}


def as_int(value: str, label: str) -> int:
    text = value.strip()
    if not text.lstrip("-").isdigit():
        raise ValueError(f"{label} must be a whole number, got '{value}'")
    return int(text)


# --------------------------------------------------------------------------- #
# Reporting
# --------------------------------------------------------------------------- #
@dataclass
class RowResult:
    row: str
    target: str
    action: str = "skipped"
    detail: str = ""
    extra: dict[str, str] = field(default_factory=dict)


def write_report(path: Path, results: list[RowResult], extra_columns: tuple[str, ...] = ()) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["row", "target", "action", *extra_columns, "detail"])
        for result in results:
            writer.writerow(
                [
                    result.row,
                    result.target,
                    result.action,
                    *[result.extra.get(column, "") for column in extra_columns],
                    result.detail,
                ]
            )


def summarise(results: list[RowResult]) -> dict[str, int]:
    tally: dict[str, int] = {}
    for result in results:
        tally[result.action] = tally.get(result.action, 0) + 1
    return tally


def configure_logging(log_path: Path, verbose: bool = False) -> None:
    console = logging.StreamHandler(sys.stdout)
    console.setLevel(logging.DEBUG if verbose else logging.INFO)
    console.setFormatter(logging.Formatter("%(message)s"))

    logging.basicConfig(
        level=logging.DEBUG,
        format="%(asctime)s  %(levelname)-8s  %(message)s",
        handlers=[logging.FileHandler(log_path, encoding="utf-8"), console],
    )


def finish(stem: Path, results: list[RowResult], extra_columns: tuple[str, ...] = ()) -> int:
    """Write the report, log a summary, and return the process exit code."""
    report_path = Path(f"{stem}_report.csv")
    write_report(report_path, results, extra_columns)
    tally = summarise(results)
    log.info("-" * 78)
    log.info(
        "Summary: %s",
        ", ".join(f"{key}={value}" for key, value in sorted(tally.items())) or "nothing processed",
    )
    log.info("Report: %s", report_path)
    return 1 if tally.get("failed") else 0
